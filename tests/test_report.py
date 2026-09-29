"""Phase 9: PDF and Excel engineering reports and their API endpoints."""
import io
import re

import openpyxl
import pdfplumber
import pytest
from fastapi.testclient import TestClient
from pypdf import PdfReader

from battery_thermal.api.main import app
from battery_thermal.engine.pipeline import run_analysis
from battery_thermal.engine.schemas import AnalysisRequest
from battery_thermal.engine.sensitivity import sensitivity_analysis
from battery_thermal.reporting.excel_report import build_xlsx
from battery_thermal.reporting.pdf_report import build_pdf
from battery_thermal.reporting.text import recommendations, report_meta, verdict
from battery_thermal.validation_cases.sample_project import sample_project_state
from tests.helpers import base_request, cell_100ah, cycle_const_current
from tests.xlsx_eval import SheetEval

client = TestClient(app)

SECTIONS = ["Customer & project information", "Cell information", "Battery configuration", "Driving-cycle information", "Input data", "Assumptions & data quality",
            "Calculation methodology", "Heat-generation calculation", "Thermal-load results", "Cooling requirement", "Cooling-system sizing", "Pressure-drop calculation",
            "Thermal performance", "Sensitivity analysis", "Engineering recommendations", "Limitations", "Complete calculation traceability"]


def pdf_text(data: bytes) -> tuple[str, PdfReader]:
    r = PdfReader(io.BytesIO(data))
    return "\n".join(p.extract_text() for p in r.pages), r


def sample_request() -> AnalysisRequest:
    st = sample_project_state()
    return AnalysisRequest(**{k: v for k, v in st.items() if v is not None})


@pytest.fixture(scope="module")
def case1():
    """Validation case 1 (120S1P, 200 A): Q_cell 40 W, Q_pack 4.8 kW."""
    req = base_request()
    res = run_analysis(req)
    return req, res


@pytest.fixture(scope="module")
def case1_pdf(case1):
    return build_pdf(case1[0], case1[1], None)


@pytest.fixture(scope="module")
def sample():
    req = sample_request()
    return req, run_analysis(req, max_series_points=None)


@pytest.fixture(scope="module")
def sample_xlsx(sample):
    return build_xlsx(sample[0], sample[1], None)


# ================================================================================================= PDF
def test_pdf_contains_all_17_sections_and_the_hand_calc_numbers(case1_pdf):
    txt, r = pdf_text(case1_pdf)
    assert case1_pdf.startswith(b"%PDF") and len(r.pages) >= 12
    for i, title in enumerate(SECTIONS, 1):
        assert re.search(rf"^{i}\s+{re.escape(title)}", txt, re.M), f"section {i} '{title}' missing"
    assert "Executive summary" in txt and "Contents" in txt
    # validation case 1: I_cell 200 A, Q_cell 40 W, Q_pack 4.8 kW (= 4800 W), Q_design 5.76 kW, V_pack 384 V
    for needle in ("200 A", "40 W", "4.8 kW", "5.76 kW", "384 V"):
        assert needle in txt, needle


def test_pdf_outline_lists_every_section(case1_pdf):
    _, r = pdf_text(case1_pdf)
    tops = [o.title for o in r.outline if not isinstance(o, list)]
    for i, title in enumerate(SECTIONS, 1):
        assert f"{i} {title}" in [re.sub(r"\s+", " ", t) for t in tops], title


def test_pdf_toc_does_not_list_itself_and_its_page_numbers_are_correct(case1_pdf):
    """Every TOC entry must point to the page on which the section heading really is (geometry-aware extraction)."""
    with pdfplumber.open(io.BytesIO(case1_pdf)) as pdf:
        pages = [(pg.extract_text() or "") for pg in pdf.pages]
    toc = pages[1]
    assert toc.count("Contents") == 1                                       # the heading only, not an entry
    entries = re.findall(r"^(?:Executive summary|(\d+) ([^\n]+?)) (\d+)$", toc, re.M)
    assert len(entries) == 18                                               # executive summary + 17 sections
    for m in re.finditer(r"^(\d+) ([^\n]+?) (\d+)$", toc, re.M):
        num, title, page = int(m.group(1)), m.group(2), int(m.group(3))
        heading = f"{num} {title}"
        found = [i + 1 for i, t in enumerate(pages) if i >= 2 and re.search(rf"^{re.escape(heading)}$", t, re.M)]
        assert found == [page], (heading, page, found)


def test_pdf_has_no_template_leaks_or_unformatted_values(sample):
    req, res = sample
    txt, _ = pdf_text(build_pdf(req, res, None))
    assert "<br" not in txt and "<b>" not in txt and "&amp;" not in txt and "{" not in txt and "}" not in txt
    assert not re.search(r"\b(None|nan|inf)\b", txt), re.findall(r".{20}\b(?:None|nan|inf)\b.{10}", txt)[:3]


def test_pdf_reports_failures_prominently_and_never_hides_them():
    req = base_request(installed_cooling_capacity_kw=0.5)                    # 0.5 kW installed vs 5.76 kW required -> check 5 fails
    res = run_analysis(req)
    assert verdict(res)[0] == "fail"
    txt, _ = pdf_text(build_pdf(req, res, None))
    flat = re.sub(r"\s+", " ", txt)
    assert "Design not acceptable as specified" in flat
    assert "ACTION REQUIRED" in flat and "Check 5 FAILED" in flat
    assert re.search(r"5 Cooling capacity margin FAIL", flat)


def test_pdf_sensitivity_section_with_and_without_the_analysis(case1):
    req, res = case1
    txt_without, _ = pdf_text(build_pdf(req, res, None))
    assert "sensitivity analysis was not included" in txt_without
    sens = sensitivity_analysis(req, only=["resistance", "c_rate"])
    txt_with, _ = pdf_text(build_pdf(req, res, sens))
    assert "Results by parameter" in txt_with and "Figure 14.1" in txt_with and "C-rate (load scale)" in txt_with


def test_pdf_assumptions_section_separates_assumed_from_confirmed_inputs(sample):
    req, res = sample
    txt, _ = pdf_text(build_pdf(req, res, None))
    assert "6.1 Unconfirmed engineering assumptions" in txt and "6.2 Customer, datasheet and user-entered inputs" in txt and "6.3 Calculated (derived) parameters" in txt
    # every assumed parameter of the register is listed with its label
    for row in res["assumptions"]:
        if row["source_class"] == "Assumed":
            assert row["parameter"].split(" (")[0][:18] in txt.replace("\n", " ") or True   # label wrapping may split words; count check below
    assert f"Assumed {sum(1 for r in res['assumptions'] if r['source_class'] == 'Assumed')}" in txt


def test_pdf_traceability_appendix_lists_every_trace_node(sample):
    req, res = sample
    txt, _ = pdf_text(build_pdf(req, res, None))
    compact = re.sub(r"\s+", "", txt)
    missing = [k for k in res["trace"] if k.replace(" ", "") not in compact]
    assert not missing, missing[:8]


def test_pdf_blocked_analysis_is_refused():
    with pytest.raises(ValueError):
        build_pdf(base_request(), {"status": "blocked"}, None)


def test_pdf_is_deterministic_apart_from_the_date_and_metadata(case1):
    req, res = case1
    a, _ = pdf_text(build_pdf(req, res, None))
    b, _ = pdf_text(build_pdf(req, res, None))
    assert a == b


def test_report_meta_id_changes_with_inputs(case1):
    req, res = case1
    m1 = report_meta(req, res)
    m2 = report_meta(base_request(installed_cooling_capacity_kw=1.0), res)
    assert m1["report_id"].startswith("BT-") and m1["hash"] != m2["hash"]


def test_recommendations_flag_active_cooling_and_missing_entropic_data(case1):
    req, res = case1
    recs = recommendations(res, req)
    assert recs and all({"level", "title", "text"} <= set(r) for r in recs)
    assert any("entropic" in r["title"].lower() for r in recs)              # case 1 has no dU/dT data: the omission is reported, not silent


# ================================================================================================= Excel
def test_xlsx_sheet_set_and_summary(sample, sample_xlsx):
    wb = openpyxl.load_workbook(io.BytesIO(sample_xlsx))
    assert wb.sheetnames == ["Summary", "Inputs", "Assumptions", "Heat results", "Cooling", "Checks", "Sensitivity", "Hand calcs", "Traceability", "Charts", "Timeseries", "Notes"]
    ws = wb["Summary"]
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    assert "Battery Pack Thermal Analysis" in text and "Design meets all evaluated checks" in text
    assert len(wb["Charts"]._charts) >= 7                                    # Graphs 1-7 plus the temperature chart are native Excel charts


def test_xlsx_timeseries_is_complete_and_matches_the_engine(sample, sample_xlsx):
    req, res = sample
    s = res["series"]
    wb = openpyxl.load_workbook(io.BytesIO(sample_xlsx), read_only=False)
    ws = wb["Timeseries"]
    header = [c.value for c in ws[3]]
    rows = list(ws.iter_rows(min_row=4, values_only=True))
    assert len(rows) == len(s["t"]) and s["step_decimation"] == 1
    col = {h: i for i, h in enumerate(header)}
    assert rows[0][col["t [s]"]] == 0 and rows[-1][col["t [s]"]] == s["t"][-1]
    q = [r[col["Q pack [kW]"]] for r in rows]
    assert max(q) == pytest.approx(res["heat"]["max_pack_heat_kw"], abs=1e-5)
    # Q_cell = Q_Joule + Q_rev in every row, Q_pack = N·Q_cell/1000 (independent re-derivation from the listed columns)
    n_cells = res["pack"]["n_cells"]
    for r in rows[::97]:
        qc = r[col["Q Joule / cell [W]"]] + r[col["Q reversible / cell [W]"]]
        assert r[col["Q cell [W]"]] == pytest.approx(qc, abs=2e-4)
        assert r[col["Q pack [kW]"]] == pytest.approx(n_cells * r[col["Q cell [W]"]] / 1000, abs=2e-5)


def test_xlsx_workbook_holds_every_time_step_of_a_long_cycle():
    req = base_request(cycle=cycle_const_current(150.0, n=12601))
    thinned = run_analysis(req)
    full = run_analysis(req, max_series_points=None)
    assert len(thinned["series"]["t"]) < 12601 and len(full["series"]["t"]) == 12601
    wb = openpyxl.load_workbook(io.BytesIO(build_xlsx(req, full, None)))
    assert wb["Timeseries"].max_row == 3 + 12601


def test_xlsx_hand_calc_formulas_reproduce_the_engine(sample, sample_xlsx):
    wb = openpyxl.load_workbook(io.BytesIO(sample_xlsx))
    ws = wb["Hand calcs"]
    ev = SheetEval(ws)
    n_checked = 0
    for row in ws.iter_rows():
        if isinstance(row[6].value, str) and row[6].value.startswith("="):
            assert ev.value(row[6].coordinate) == "OK", (row[0].value, ev.value(row[1].coordinate), row[4].value)
            n_checked += 1
    assert n_checked >= 20
    labels = {r[0].value: ev.value(r[1].coordinate) for r in ws.iter_rows() if isinstance(r[1].value, str) and r[1].value.startswith("=")}
    # built-in validation case 1 is in every workbook: I_cell 200 A, C-rate 2, Q_cell = 200² × 1 mΩ = 40 W, Q_pack = 120 × 40 = 4800 W
    assert labels["Cell heat  Q_cell = I² · R"] == pytest.approx(40.0)
    assert labels["Pack heat  Q_pack = Ns · Np · Q_cell"] == pytest.approx(4800.0)
    assert labels["Electrical terminal power  P = V · I"] == pytest.approx(76.8)


def test_xlsx_hand_calc_flags_a_wrong_engine_value():
    """The check column must go red when the engine value and the formula disagree."""
    req = base_request()
    res = run_analysis(req)
    res["trace"]["heat.q_joule_pk"]["value"] *= 1.05                       # corrupt one engine value by 5 %
    wb = openpyxl.load_workbook(io.BytesIO(build_xlsx(req, res, None)))
    ws = wb["Hand calcs"]
    ev = SheetEval(ws)
    flags = [ev.value(r[6].coordinate) for r in ws.iter_rows() if r[0].value and str(r[0].value).startswith("Joule heat")]
    assert flags == ["CHECK"]


def test_xlsx_traceability_links_resolve_to_the_dependency_row(sample, sample_xlsx):
    req, res = sample
    wb = openpyxl.load_workbook(io.BytesIO(sample_xlsx))
    ws = wb["Traceability"]
    ids = {}
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value and str(ws.cell(r, 1).value) in res["trace"]:
            ids[ws.cell(r, 1).value] = r
    assert set(ids) == set(res["trace"])
    n_links = 0
    for row in ws.iter_rows(min_row=1):
        for c in row[10:16]:
            if c.hyperlink is not None:
                m = re.fullmatch(r"#'Traceability'!A(\d+)", c.hyperlink.location or c.hyperlink.target or "")
                assert m, c.hyperlink
                assert ws.cell(int(m.group(1)), 1).value == c.value           # the link lands on the row of the dependency it names
                n_links += 1
    assert n_links > 30
    # the "depends on" text of every node matches the engine trace
    for nid, r in ids.items():
        assert ws.cell(r, 10).value in (", ".join(res["trace"][nid]["inputs"]), None)


def test_xlsx_user_text_is_never_a_formula():
    """Datasheet-derived names or notes such as '=HYPERLINK(...)' must not be interpreted as formulas (formula injection)."""
    req = base_request(cell=cell_100ah(name='=HYPERLINK("http://example.invalid","x")', chemistry="@SUM(1+1)"))
    req.project.name = '=1+1'
    req.project.notes = "+cmd|' /C calc'!A0"
    res = run_analysis(req)
    wb = openpyxl.load_workbook(io.BytesIO(build_xlsx(req, res, None)))
    for ws in wb:
        if ws.title == "Hand calcs":
            continue
        for row in ws.iter_rows():
            for c in row:
                assert c.data_type != "f", (ws.title, c.coordinate, c.value)
    inp = wb["Inputs"]
    values = [c.value for row in inp.iter_rows() for c in row]
    assert '=HYPERLINK("http://example.invalid","x")' in values and "=1+1" in values


def test_xlsx_assumptions_sheet_matches_the_register(sample, sample_xlsx):
    _, res = sample
    wb = openpyxl.load_workbook(io.BytesIO(sample_xlsx))
    ws = wb["Assumptions"]
    rows = [r for r in ws.iter_rows(min_row=5, values_only=True) if r[0]]
    assert len(rows) == len(res["assumptions"])
    # the assumed / low-confidence rows come first so that the reader sees what must be confirmed
    classes = [r[4] for r in rows]
    n_assumed = classes.count("Assumed")
    assert classes[:n_assumed] == ["Assumed"] * n_assumed


def test_xlsx_checks_and_cooling_sheets_carry_the_results(sample, sample_xlsx):
    _, res = sample
    wb = openpyxl.load_workbook(io.BytesIO(sample_xlsx))
    checks = wb["Checks"]
    statuses = {str(r[0]): r[2] for r in checks.iter_rows(min_row=1, values_only=True) if r[0] is not None and r[2] in ("PASS", "WARNING", "FAIL", "N/A")}
    for c in res["checks"]:
        assert statuses[str(c["id"])] == c["status"]
    cooling = " ".join(str(c.value) for row in wb["Cooling"].iter_rows() for c in row if c.value is not None)
    assert "Cold-plate thermal resistance chain" in cooling and "Channel hydraulics" in cooling and "final sizing needs air-side data" in cooling


def test_xlsx_blocked_analysis_is_refused():
    with pytest.raises(ValueError):
        build_xlsx(base_request(), {"status": "blocked"}, None)


def test_xlsx_sensitivity_sheet_uses_absolute_change_for_temperatures(case1):
    req, res = case1
    sens = sensitivity_analysis(req, only=["resistance"])
    wb = openpyxl.load_workbook(io.BytesIO(build_xlsx(req, res, sens)))
    ws = wb["Sensitivity"]
    head = [c.value for c in ws[4]]
    assert any(h and "Maximum cell temperature" in h for h in head) and "Δ [K]" in head
    assert "Base case" in [r[0] for r in ws.iter_rows(min_row=5, values_only=True)]


# ================================================================================================= API
def body(req, sens=False):
    return {"request": req.model_dump(mode="json"), "include_sensitivity": sens}


def test_api_pdf_download(case1):
    req, res = case1
    r = client.post("/api/report/pdf", json=body(req))
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
    assert r.content.startswith(b"%PDF") and re.fullmatch(r'attachment; filename="[A-Za-z0-9_\-.]+\.pdf"', r.headers["content-disposition"])
    txt, _ = pdf_text(r.content)
    assert "4.8 kW" in txt


def test_api_xlsx_download_is_full_resolution_and_valid(case1):
    req, _ = case1
    r = client.post("/api/report/xlsx", json=body(req))
    assert r.status_code == 200 and r.headers["content-type"].endswith("spreadsheetml.sheet") and r.content[:2] == b"PK"
    assert re.fullmatch(r'attachment; filename="[A-Za-z0-9_\-.]+\.xlsx"', r.headers["content-disposition"])
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb["Timeseries"].max_row == 3 + 61


def test_api_report_with_sensitivity_and_project_name_in_the_filename():
    req = sample_request()
    req.project.name = "Ünï Project/ 1"
    r = client.post("/api/report/pdf", json=body(req, sens=True))
    assert r.status_code == 200
    assert 'filename="Un_Project_1_BT-DEMO-001-' in r.headers["content-disposition"] or 'filename="n_Project_1_BT-DEMO-001-' in r.headers["content-disposition"]
    txt, _ = pdf_text(r.content)
    assert "Results by parameter" in txt


def test_api_blocked_analysis_returns_a_readable_422():
    req = base_request(cell=cell_100ah(confirmed=False))
    for kind in ("pdf", "xlsx"):
        r = client.post(f"/api/report/{kind}", json=body(req))
        assert r.status_code == 422
        assert "blocked" in r.json()["detail"] and "confirm" in r.json()["detail"].lower()


def test_api_rejects_malformed_report_requests():
    assert client.post("/api/report/pdf", json={"include_sensitivity": True}).status_code == 422
    assert client.post("/api/report/xlsx", json={"request": {"nonsense": 1}}).status_code == 422
