"""Generate the sample data files (XLSX / PDF datasheets, driving cycles).

Everything generated here is SYNTHETIC test data and is labelled as such in the files themselves.
Run from the repository root:  python scripts/make_sample_data.py
"""
from __future__ import annotations

import math
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
OUT = ROOT / "sample_data"

from battery_thermal.ingestion.common import read_csv_rows, to_float          # noqa: E402
from battery_thermal.ingestion.datasheet import split_blocks                   # noqa: E402


def _num(c: str):
    v = to_float(c)
    return v if v is not None else c


def make_cell_xlsx() -> Path:
    from openpyxl import Workbook
    rows = read_csv_rows((OUT / "cell_datasheet_LFP100Ah.csv").read_bytes())
    wb = Workbook()
    ws = wb.active
    ws.title = "Cell"
    ws.append(["# SYNTHETIC ILLUSTRATIVE DATA for software testing only - not a real product datasheet"])
    for title, block in split_blocks(rows):
        target = ws if not title else wb.create_sheet(title.replace(" ", "_")[:28])
        for r in block:
            target.append([_num(c) for c in r])
    # full-featured extras that the CSV omits: entropic coefficient table and a 2-D resistance map
    d = wb.create_sheet("dUdT_vs_SOC")
    d.append(["SOC (%)", "dU/dT (mV/K)"])
    for soc, v in [(0, 0.10), (10, 0.05), (20, -0.02), (40, -0.05), (60, -0.03), (80, 0.02), (100, 0.08)]:
        d.append([soc, v])
    m = wb.create_sheet("R_map")
    m.append(["SOC (%) \\ T (°C) - R_DC (mΩ)", -10, 0, 10, 25, 40, 55])
    soc_f = {0: 0.85, 10: 0.66, 20: 0.58, 50: 0.52, 80: 0.53, 100: 0.62}
    t_f = {-10: 1.85, 0: 1.15, 10: 0.78, 25: 0.52, 40: 0.41, 55: 0.36}
    for soc, fs in soc_f.items():
        m.append([soc] + [round(fs / 0.52 * ft, 4) for ft in t_f.values()])
    path = OUT / "cell_datasheet_LFP100Ah.xlsx"
    wb.save(path)
    return path


def make_cell_pdf() -> Path:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    st = getSampleStyleSheet()
    path = OUT / "cell_datasheet_LFP100Ah.pdf"
    doc = SimpleDocTemplate(str(path), pagesize=A4)
    rows = [["Item", "Specification"],
            ["Nominal Capacity", "100 Ah"], ["Nominal Voltage", "3.2 V"], ["Charge Cut-off Voltage", "3.65 V"],
            ["Discharge Cut-off Voltage", "2.5 V"], ["DC Internal Resistance", "≤ 0.6 mΩ"],
            ["AC Impedance (1 kHz)", "0.28 mΩ"], ["Max continuous discharge current", "100 A"],
            ["Max continuous charge current", "50 A"], ["Weight", "2.05 kg"],
            ["Dimensions (T x W x H)", "71 x 174 x 207 mm"],
            ["Operating temperature (discharge)", "-20 ~ 60 °C"], ["Operating temperature (charge)", "0 ~ 55 °C"]]
    t = Table(rows, colWidths=[260, 200])
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey)]))
    doc.build([Paragraph("SYN-LFP-100Ah prismatic LFP cell - product specification", st["Title"]),
               Paragraph("SYNTHETIC ILLUSTRATIVE DATA for software testing only - not a real product.", st["Normal"]),
               Spacer(1, 12), t, Spacer(1, 12),
               Paragraph("Chemistry: LiFePO4 (LFP)  Cell type: prismatic", st["Normal"]),
               Paragraph("Pulse discharge current 200 A (10 s)", st["Normal"])])
    return path


# ---------------------------------------------------------------------------------------------
# driving cycles (synthetic)
# ---------------------------------------------------------------------------------------------
def _speed_profile(dt: float = 1.0) -> list[float]:
    """Synthetic urban + extra-urban + motorway speed trace [km/h], ~1200 s. Not a standardised cycle."""
    knots = [(0, 0), (10, 0), (25, 30), (40, 30), (50, 0), (60, 0), (80, 45), (110, 50), (125, 20), (135, 0), (145, 0),
             (170, 50), (200, 60), (240, 60), (260, 30), (270, 0), (285, 0), (330, 80), (400, 90), (480, 100), (560, 100),
             (620, 70), (660, 90), (740, 110), (820, 120), (900, 120), (960, 80), (1010, 30), (1050, 0), (1080, 0),
             (1110, 40), (1150, 50), (1180, 0), (1200, 0)]
    out, t = [], 0.0
    while t <= 1200 + 1e-9:
        for (t0, v0), (t1, v1) in zip(knots, knots[1:]):
            if t0 <= t <= t1:
                out.append(v0 + (v1 - v0) * (t - t0) / (t1 - t0))
                break
        t += dt
    return out


def make_drive_cycles() -> list[Path]:
    import csv
    v = _speed_profile()
    t = list(range(len(v)))
    p1 = OUT / "drive_cycle_speed_only.csv"
    with p1.open("w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["# SYNTHETIC speed trace for software testing - not a standardised cycle"])
        w.writerow(["Time [s]", "Vehicle speed [km/h]"])
        for ti, vi in zip(t, v):
            w.writerow([ti, round(vi, 3)])
    # battery current/power/SOC file generated from a simple road-load model (documented, reproducible)
    m, crr, cd, area, eta, aux = 1800.0, 0.009, 0.28, 2.2, 0.90, 500.0
    g, rho = 9.80665, 1.225
    n = len(v)
    vms = [x / 3.6 for x in v]
    acc = [(vms[min(i + 1, n - 1)] - vms[max(i - 1, 0)]) / (t[min(i + 1, n - 1)] - t[max(i - 1, 0)]) for i in range(n)]
    p2 = OUT / "drive_cycle_battery.csv"
    soc, cap_j = 90.0, 38.4e3 * 3600
    with p2.open("w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["# SYNTHETIC battery power/current trace (road-load model, 1800 kg, eta 0.90, aux 0.5 kW) - test data only"])
        w.writerow(["Time [s]", "Speed [km/h]", "Acceleration [m/s2]", "Battery power [kW]", "Battery current [A]", "SOC [%]"])
        for i in range(n):
            force = m * acc[i] + (crr * m * g if vms[i] > 0.01 else 0) + 0.5 * rho * cd * area * vms[i] ** 2
            pw = force * vms[i]
            pb = (pw / eta if pw >= 0 else pw * eta) + aux
            i_pack = pb / 384.0 * 1000.0
            w.writerow([t[i], round(v[i], 3), round(acc[i], 4), round(pb / 1000.0, 4), round(i_pack, 3), round(soc, 4)])
            soc -= pb * 1.0 / cap_j * 100
    return [p1, p2]


if __name__ == "__main__":
    print(make_cell_xlsx())
    print(make_cell_pdf())
    for p in make_drive_cycles():
        print(p)
