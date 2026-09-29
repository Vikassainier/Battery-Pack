"""Cell datasheet parser (CSV / Excel / PDF).

The parser only *proposes* values. Every extracted item carries its unit conversion, a confidence
level, and the source text/location so the user can review, edit and confirm before any calculation.
Nothing here is trusted blindly and nothing is invented: missing required parameters are listed with an
optional, clearly-labelled engineering suggestion the user must explicitly accept.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field, asdict
from typing import Any

from ..engine.validation import Issue, ERROR, WARNING, INFO
from .common import IngestionError, SheetRows, file_kind, read_sheets, to_float

# ---------------------------------------------------------------------------------------------
# result containers
# ---------------------------------------------------------------------------------------------
@dataclass
class ExtractedField:
    path: str
    value: Any
    unit: str = ""
    confidence: str = "medium"          # high | medium | low
    snippet: str = ""
    location: str = ""
    note: str = ""


@dataclass
class Suggestion:
    path: str
    value: Any
    unit: str
    confidence: str
    rationale: str


@dataclass
class DatasheetResult:
    source_name: str
    source_type: str
    fields: dict[str, ExtractedField] = field(default_factory=dict)
    curves: dict[str, dict] = field(default_factory=dict)
    maps: dict[str, dict] = field(default_factory=dict)
    issues: list[Issue] = field(default_factory=list)
    missing: list[dict] = field(default_factory=list)
    suggestions: dict[str, Suggestion] = field(default_factory=dict)

    def cell_proposal(self) -> dict:
        """CellSpec-compatible dict of the *proposed* values (``confirmed`` is always False)."""
        cell: dict[str, Any] = {}
        for path, f in self.fields.items():
            if path.startswith("cell."):
                cell[path[5:]] = f.value
        for key, c in self.curves.items():
            cell[key] = {"x": c["x"], "y": c["y"], "x_label": c.get("x_label", ""), "y_label": c.get("y_label", "")}
        for key, m in self.maps.items():
            cell[key] = {"x": m["x"], "y": m["y"], "z": m["z"], "x_label": "SOC [%]", "y_label": "T [degC]",
                         "z_label": m.get("z_label", "")}
        cell["confirmed"] = False
        return cell

    def to_dict(self) -> dict:
        return {
            "source_name": self.source_name, "source_type": self.source_type,
            "cell": self.cell_proposal(),
            "fields": {p: asdict(f) for p, f in self.fields.items()},
            "curves": self.curves, "maps": self.maps,
            "missing": self.missing,
            "suggestions": {p: asdict(s) for p, s in self.suggestions.items()},
            "issues": [i.to_dict() for i in self.issues],
        }


# ---------------------------------------------------------------------------------------------
# units
# ---------------------------------------------------------------------------------------------
def canon_unit(u: str) -> str:
    u = (u or "").strip().lower()
    for a, b in (("µ", "u"), ("μ", "u"), ("ω", "ohm"), ("Ω".lower(), "ohm"), ("°", ""), ("º", ""), ("·", ""),
                 ("(", ""), (")", ""), (" ", ""), (".", ""), ("*", ""), ("⋅", ""), ("^", "")):
        u = u.replace(a, b)
    return u


UNITS: dict[str, dict[str, Any]] = {
    "capacity": {"ah": 1.0, "mah": 1e-3},
    "voltage": {"v": 1.0, "mv": 1e-3, "vdc": 1.0},
    "resistance": {"mohm": 1.0, "ohm": 1e3, "uohm": 1e-3, "milliohm": 1.0, "mohms": 1.0},
    "mass": {"kg": 1.0, "g": 1e-3, "kgs": 1.0},
    "length": {"mm": 1.0, "cm": 10.0, "m": 1000.0},
    "temp": {"c": 1.0, "degc": 1.0, "celsius": 1.0, "k": "K"},
    "cp": {"j/kgk": 1.0, "j/kgc": 1.0, "kj/kgk": 1000.0, "j/gk": 1000.0, "jkg-1k-1": 1.0},
    "dudt": {"mv/k": 1.0, "v/k": 1000.0, "uv/k": 1e-3},
    "crate": {"c": 1.0, "a": "A"},
    "time": {"s": 1.0, "sec": 1.0, "min": 60.0, "ms": 1e-3},
    "none": {},
}
DEFAULT_UNIT = {"capacity": "Ah", "voltage": "V", "resistance": "mΩ", "mass": "kg", "length": "mm", "temp": "°C",
                "cp": "J/(kg·K)", "dudt": "mV/K", "crate": "C", "time": "s", "none": ""}
# plausible ranges after conversion: below/above -> confidence downgraded with a note
PLAUSIBLE = {
    "cell.capacity_ah": (0.1, 1500), "cell.v_nom": (1.0, 4.6), "cell.v_max": (1.5, 5.0), "cell.v_min": (0.5, 4.0),
    "cell.r_dc_mohm": (0.005, 200), "cell.r_ac_mohm": (0.005, 200), "cell.mass_kg": (0.005, 30),
    "cell.max_charge_c": (0.05, 25), "cell.max_discharge_c": (0.05, 40), "cell.pulse_discharge_c": (0.05, 60),
    "cell.pulse_charge_c": (0.05, 40), "cell.cp_j_kg_k": (300, 3000),
}


def convert(value: float, unit_kind: str, unit_txt: str) -> tuple[float | None, str, bool]:
    """Return (converted value, canonical unit label, unit_was_assumed)."""
    table = UNITS[unit_kind]
    cu = canon_unit(unit_txt)
    if unit_kind == "none":
        return value, "", False
    if unit_kind == "resistance" and cu in ("mw", "w"):
        # PDF fonts frequently lose the Ω glyph so "mΩ" is extracted as "mW"; a resistance can never be a power
        return value * (1.0 if cu == "mw" else 1e3), "mΩ", True
    if cu in table:
        f = table[cu]
        if f == "K":
            return value - 273.15, "°C", False
        if f == "A":
            return value, "A", False
        return value * f, DEFAULT_UNIT[unit_kind], False
    return value, DEFAULT_UNIT[unit_kind], True


# ---------------------------------------------------------------------------------------------
# value parsing
# ---------------------------------------------------------------------------------------------
_NUM = r"[-+−–]?\s?\d+(?:[.,]\d+)?"
VALUE_RE = re.compile(
    rf"(?P<q>[<>≤≥~≈£³]=?|max\.?|min\.?|typ\.?|approx\.?)?\s*(?P<n1>{_NUM})"
    rf"(?:\s*(?:~|to|…|\.\.|[-–—])\s*(?P<n2>{_NUM}))?\s*(?P<u>[^\s\d,;:]{{1,14}})?", re.I)


@dataclass
class ParsedValue:
    n1: float
    n2: float | None
    unit: str
    qualifier: str
    raw: str


def parse_value(text: str) -> ParsedValue | None:
    t = (text or "").replace(" ", " ").strip()
    m = VALUE_RE.search(t)
    if not m:
        return None
    n1 = to_float(m.group("n1").replace(" ", ""))
    if n1 is None:
        return None
    n2 = to_float(m.group("n2").replace(" ", "")) if m.group("n2") else None
    unit = (m.group("u") or "").strip(".,;)")
    if unit.lower() in ("to", "~", "typ", "max", "min", "at", "@", "x", "×"):
        unit = ""
    q = (m.group("q") or "").lower()
    q = {"£": "≤", "³": "≥"}.get(q.rstrip("="), q)           # Symbol-font glyph fallbacks seen in PDF text extraction
    return ParsedValue(n1, n2, unit, q, m.group(0).strip())


# ---------------------------------------------------------------------------------------------
# field catalogue: label pattern -> canonical path
# ---------------------------------------------------------------------------------------------
@dataclass(frozen=True)
class FieldDef:
    path: str
    pattern: str
    unit_kind: str
    exclude: str | None = None
    kind: str = "num"           # num | range | text | dims | crate
    note: str = ""
    conf_cap: str | None = None


FIELD_DEFS: list[FieldDef] = [
    FieldDef("cell.r_ac_mohm", r"\b(ac|1\s?khz)\b.*(resist|imped)|imped|\bacir\b|\bacr\b|resist.*1\s?khz", "resistance"),
    FieldDef("cell.r_dc_mohm", r"\b(dc|hppc|pulse)\b.*resist|\bdcir\b|\bdcr\b|resist.*\bdc\b", "resistance"),
    FieldDef("cell.r_ac_mohm", r"internal resistance|\bresistance\b", "resistance", exclude=r"\bvs\b|versus|temperature",
             note="Resistance without an AC/DC qualifier: treated as AC impedance (not used for heat). "
                  "Enter the DC resistance (DCIR) if known.", conf_cap="low"),
    FieldDef("cell.cp_j_kg_k", r"specific\s*heat|heat\s*capacity|thermal\s*capacity|\bcp\b", "cp"),
    FieldDef("cell.capacity_ah", r"(?<![\w])(nominal|rated|typical|typ|minimum|min)?\s*(cell\s+)?capacity", "capacity",
             exclude=r"retention|\bvs\b|versus|loss|fade|temperature|specific|heat|thermal|energy|@\s*-|at\s*-"),
    FieldDef("cell.v_nom", r"(nominal|rated|average|mean|operating|working)\s+(cell\s+)?voltage|^voltage$", "voltage"),
    FieldDef("cell.v_max", r"(max(imum)?|upper[\s-]*limit|\bcharge\b[\s-]*cut.?off|\bcharging\b[\s-]*cut.?off|end[\s-]of[\s-]charge|"
                           r"\bcharge\b|\bcharging\b)\s*(cut.?off\s*)?voltage", "voltage", exclude=r"\bdischarg"),
    FieldDef("cell.v_min", r"(min(imum)?|lower[\s-]*limit|discharge[\s-]*cut.?off|discharging[\s-]*cut.?off|end[\s-]of[\s-]discharge|"
                           r"cut.?off)\s*(cut.?off\s*)?voltage", "voltage"),
    FieldDef("cell.max_charge_c", r"(max(imum)?|continuous|standard|recommended)?\s*(continuous\s*)?\bcharg(e|ing)\b\s*(current|rate|c-rate)",
             "crate", exclude=r"pulse|peak|discharg|cut.?off|voltage", kind="crate"),
    FieldDef("cell.max_discharge_c", r"(max(imum)?|continuous|standard|recommended)?\s*(continuous\s*)?discharg(e|ing)\s*(current|rate|c-rate)",
             "crate", exclude=r"pulse|peak|cut.?off|voltage", kind="crate"),
    FieldDef("cell.pulse_charge_c", r"(pulse|peak)\s*charg(e|ing)\s*(current|rate|c-rate)?|charg(e|ing)\s*(pulse|peak)",
             "crate", kind="crate"),
    FieldDef("cell.pulse_discharge_c", r"(pulse|peak)\s*discharg(e|ing)\s*(current|rate|c-rate)?|discharg(e|ing)\s*(pulse|peak)|"
                                       r"(max(imum)?\s*)?(pulse|peak)\s*(current|rate)", "crate", kind="crate"),
    FieldDef("cell.pulse_duration_s", r"pulse\s*(duration|time|length)|peak\s*(duration|time)", "time"),
    FieldDef("cell.mass_kg", r"\b(weight|mass)\b", "mass", exclude=r"energy density|specific"),
    FieldDef("entropic.constant_mv_per_k", r"entropic|du\s?/\s?dt|dudt|temperature\s*coefficient\s*of\s*(ocv|voltage)", "dudt",
             conf_cap="low", note="Single dU/dT value - the SOC dependence is not captured."),
    FieldDef("cell.__dims__", r"dimension|\bsize\b", "length", kind="dims"),
    FieldDef("cell.length_mm", r"\blength\b", "length"),
    FieldDef("cell.width_mm", r"\bwidth\b", "length"),
    FieldDef("cell.width_mm", r"\bthickness\b", "length", note="Thickness mapped to 'width' - verify axes."),
    FieldDef("cell.height_mm", r"\bheight\b", "length"),
    FieldDef("cell.diameter_mm", r"\bdiameter\b", "length"),
    FieldDef("cell.__trange_rec__", r"(recommended|optimal|ideal|best)\s*(operating\s*)?temp", "temp", kind="range"),
    FieldDef("cell.__trange_charge__", r"(charg(e|ing))\s*(operating\s*)?temp|temp.*\bcharg(e|ing)\b", "temp", kind="range",
             exclude=r"discharg"),
    FieldDef("cell.__trange_op__", r"(discharg(e|ing)\s*)?(operating|working|ambient)\s*temp|temp.*discharg|temperature\s*range",
             "temp", kind="range", exclude=r"storage|recommended|optimal"),
    FieldDef("cell.chemistry", r"chemistry|cathode|electrochem|cell\s*type|type\s*of\s*cell|\bsystem\b", "none", kind="text"),
    FieldDef("cell.form_factor", r"\bform\b|\bformat\b|\bshape\b|construction|package|\bcase\b", "none", kind="text"),
    FieldDef("cell.name", r"model|part\s*(no|number)|product|cell\s*name|\btype\b", "none", kind="text"),
]
_COMPILED = [(d, re.compile(d.pattern, re.I), re.compile(d.exclude, re.I) if d.exclude else None) for d in FIELD_DEFS]

CHEMISTRY = [("LFP", r"lfp|lifepo|iron\s*phosphate"), ("NMC", r"nmc|ncm|nickel.?manganese.?cobalt"),
             ("NCA", r"\bnca\b|nickel.?cobalt.?alum"), ("LTO", r"\blto\b|titanate"), ("LCO", r"\blco\b|lithium cobalt oxide"),
             ("LMO", r"\blmo\b|manganese oxide"), ("Na-ion", r"na-?ion|sodium")]
FORMS = [("cylindrical", r"cylind|\b(14500|18650|20700|21700|26650|32140|4680|4695)\b"), ("prismatic", r"prismatic|\bprism"),
         ("pouch", r"pouch|soft.?pack|polymer|laminate")]


def _norm_label(label: str) -> str:
    # keep bracketed words that qualify the parameter ("(charge)", "(DCIR)"), drop unit/condition brackets ("(mΩ)", "(1 kHz)")
    s = re.sub(r"[\[(]([^\])]*)[\])]", lambda m: f" {m.group(1)} " if re.search(r"[a-z]{4,}", m.group(1).lower()) else " ",
               label.lower())
    s = s.replace("_", " ").replace("/", " / ")
    return re.sub(r"\s+", " ", s).strip(" :.-")


def match_field(label: str) -> FieldDef | None:
    norm = _norm_label(label)
    if not norm:
        return None
    for d, rx, ex in _COMPILED:
        if rx.search(norm) and not (ex and ex.search(norm)):
            return d
    return None


# ---------------------------------------------------------------------------------------------
# extraction engine
# ---------------------------------------------------------------------------------------------
class _Extractor:
    def __init__(self, name: str, source_type: str):
        self.res = DatasheetResult(name, source_type)
        self.base_conf = {"csv": "high", "xlsx": "high", "pdf": "medium"}.get(source_type, "medium")
        self.pending_amp: dict[str, tuple[float, str, str]] = {}      # path -> (amps, snippet, location)
        self.cap_curve_ah: dict | None = None

    # -- helpers ------------------------------------------------------------------------------
    def issue(self, code, sev, fld, msg, hint=""):
        self.res.issues.append(Issue(code, sev, fld, msg, hint))

    def add(self, path, value, unit, conf, snippet, location, note=""):
        lo_hi = PLAUSIBLE.get(path)
        if lo_hi and isinstance(value, (int, float)) and not (lo_hi[0] <= value <= lo_hi[1]):
            conf = "low"
            note = (note + " " if note else "") + f"Value {value:g} {unit} is outside the typical range {lo_hi[0]:g}-{lo_hi[1]:g}: check units."
            self.issue("DS_OUT_OF_TYPICAL_RANGE", WARNING, path, f"{path}: {value:g} {unit} looks atypical - verify against the datasheet.")
        rank = {"high": 3, "medium": 2, "low": 1}
        old = self.res.fields.get(path)
        if old and rank[old.confidence] >= rank[conf]:
            return
        self.res.fields[path] = ExtractedField(path, value, unit, conf, snippet[:160], location, note)

    def _conf(self, unit_assumed: bool, cap: str | None) -> str:
        c = self.base_conf
        if unit_assumed:
            c = "medium" if c == "high" else "low"
        if cap:
            c = cap if {"high": 3, "medium": 2, "low": 1}[cap] < {"high": 3, "medium": 2, "low": 1}[c] else c
        return c

    # -- one label/value pair ---------------------------------------------------------------------
    def handle_pair(self, label: str, value_cells: list[str], location: str, context: str = "") -> None:
        d = match_field(label)
        if d is None:
            return
        snippet = f"{label}: {' '.join(c for c in value_cells if c)}".strip()
        joined = " ".join(c for c in value_cells if c)
        if d.kind == "text":
            self._handle_text(d, label, joined, snippet, location)
            return
        pv = None
        vcell_idx = -1
        for i, c in enumerate(value_cells):
            pv = parse_value(c)
            if pv:
                vcell_idx = i
                break
        if pv is None:
            return
        unit_txt = pv.unit
        if not unit_txt:
            for c in value_cells[vcell_idx + 1:]:
                if c and parse_value(c) is None and len(c) <= 14:
                    unit_txt = c
                    break
        if not unit_txt:
            m = re.search(r"[\[(]([^\])]{1,12})[\])]", label)
            if m and canon_unit(m.group(1)) in UNITS[d.unit_kind]:
                unit_txt = m.group(1)
        if d.kind == "dims":
            self._handle_dims(d, label, joined, snippet, location)
            return
        if d.kind == "range":
            self._handle_range(d, label, pv, unit_txt, snippet, location)
            return
        self._handle_numeric(d, label, pv, unit_txt, snippet, location, joined + " " + label + " " + context)

    def _handle_text(self, d: FieldDef, label: str, text: str, snippet: str, location: str) -> None:
        text = text.strip()
        if not text:
            return
        conf = self._conf(False, "low" if d.path == "cell.name" else None)     # derived from a part number -> low
        for name, rx in CHEMISTRY:
            if re.search(rx, text, re.I):
                self.add("cell.chemistry", name, "", conf, snippet, location)
                break
        for name, rx in FORMS:
            if re.search(rx, text, re.I):
                self.add("cell.form_factor", name, "", conf, snippet, location)
                break
        if d.path == "cell.name":
            self.add("cell.name", text[:60], "", self._conf(False, None), snippet, location)

    def _handle_dims(self, d, label, joined, snippet, location) -> None:
        m = re.search(r"(\d+(?:[.,]\d+)?)\s*[x×*]\s*(\d+(?:[.,]\d+)?)(?:\s*[x×*]\s*(\d+(?:[.,]\d+)?))?\s*(mm|cm|m)?", joined, re.I)
        if not m:
            return
        nums = [to_float(g) for g in m.groups()[:3] if g]
        unit = m.group(4) or (re.search(r"\b(mm|cm|m)\b", label, re.I).group(1) if re.search(r"\b(mm|cm|m)\b", label, re.I) else "")
        letters = re.findall(r"\b([lwhtd])\b", label.lower()) if re.search(r"[x×]", label.lower()) else []
        if len(letters) == len(nums):
            mp = {"l": "cell.length_mm", "w": "cell.length_mm" if "t" in letters else "cell.width_mm",
                  "t": "cell.width_mm", "h": "cell.height_mm", "d": "cell.diameter_mm"}
            paths = [mp[x] for x in letters]
        else:
            paths = ["cell.length_mm", "cell.width_mm", "cell.height_mm"][:len(nums)]
        for p, v in zip(paths, nums):
            cv, u, assumed = convert(v, "length", unit)
            self.add(p, cv, "mm", self._conf(assumed, "medium"), snippet, location,
                     "Axis order as printed in the datasheet - verify which dimension is length/width/height.")

    def _handle_range(self, d, label, pv, unit_txt, snippet, location) -> None:
        lo, hi = pv.n1, pv.n2
        if hi is None:
            return
        lo_c, _, a1 = convert(lo, "temp", unit_txt or "C")
        hi_c, _, _ = convert(hi, "temp", unit_txt or "C")
        conf = self._conf(not unit_txt, None)
        pre = {"cell.__trange_rec__": ("t_rec_min_c", "t_rec_max_c"), "cell.__trange_charge__": ("t_charge_min_c", "t_charge_max_c"),
               "cell.__trange_op__": ("t_op_min_c", "t_op_max_c")}[d.path]
        self.add(f"cell.{pre[0]}", min(lo_c, hi_c), "°C", conf, snippet, location)
        self.add(f"cell.{pre[1]}", max(lo_c, hi_c), "°C", conf, snippet, location)

    def _handle_numeric(self, d: FieldDef, label: str, pv: ParsedValue, unit_txt: str, snippet: str, location: str,
                        context: str) -> None:
        path = d.path
        val = pv.n1
        note = d.note
        if pv.qualifier in ("<", "≤", "max", "max.", "<="):
            note = (note + " " if note else "") + "Datasheet gives a maximum (≤) - the typical value may be lower."
        elif pv.qualifier in (">", "≥", "min", "min.", ">="):
            note = (note + " " if note else "") + "Datasheet gives a minimum (≥)."
        if pv.n2 is not None and d.kind != "range":
            note = (note + " " if note else "") + f"A range was given ({pv.n1:g} to {pv.n2:g}); the first value is proposed."
        if path == "cell.capacity_ah" and re.search(r"\bmin", _norm_label(label)):
            note = (note + " " if note else "") + "Minimum (not nominal) capacity."
        if path == "cell.v_nom" and pv.n2 is not None:                 # "operating voltage 2.5-3.65 V" is a window, not the nominal
            lo, hi = sorted((pv.n1, pv.n2))
            cv_lo, _, _ = convert(lo, "voltage", unit_txt or "V")
            cv_hi, _, _ = convert(hi, "voltage", unit_txt or "V")
            self.add("cell.v_min", cv_lo, "V", "low", snippet, location, "Lower end of an operating-voltage window.")
            self.add("cell.v_max", cv_hi, "V", "low", snippet, location, "Upper end of an operating-voltage window.")
            return
        if d.kind == "crate":
            cu = canon_unit(unit_txt)
            if cu == "a":
                self.pending_amp[path] = (val, snippet, location)
                return
            v, u, assumed = convert(val, "crate", unit_txt or "C")
            if re.search(r"\bstandard\b|\brecommended\b", _norm_label(label)):
                note = (note + " " if note else "") + "'Standard/recommended' rate - not necessarily the maximum capability."
            self.add(path, v, "C", self._conf(assumed and not re.search(r"\d\s*c\b", pv.raw, re.I), None), snippet, location, note)
            dm = re.search(r"(\d+(?:\.\d+)?)\s*(s|sec)\b", context, re.I)
            if dm and path in ("cell.pulse_discharge_c", "cell.pulse_charge_c"):
                self.add("cell.pulse_duration_s", float(dm.group(1)), "s", self._conf(False, "medium"), snippet, location)
            return
        v, u, assumed = convert(val, d.unit_kind, unit_txt)
        if d.unit_kind == "mass" and assumed and v > 20:
            v, u = v / 1000.0, "kg"
            note = (note + " " if note else "") + "No unit given: value looked like grams, converted to kg - verify."
            assumed = True
        if assumed:
            glyph = d.unit_kind == "resistance" and canon_unit(unit_txt) in ("mw", "w")
            why = (f"Unit '{unit_txt}' read as {DEFAULT_UNIT[d.unit_kind]} (the Ω glyph is often lost when text is extracted from a PDF)."
                   if glyph else f"No/unknown unit '{unit_txt}': assumed {DEFAULT_UNIT[d.unit_kind]}.")
            note = (note + " " if note else "") + why
            self.issue("DS_UNIT_ASSUMED", WARNING, path, f"{path}: {why}")
        if path in ("cell.r_ac_mohm", "cell.r_dc_mohm"):
            tm = re.search(r"(?:@|at)\s*(-?\d+(?:\.\d+)?)\s*°?\s*c\b", context, re.I)
            sm = re.search(r"(\d+(?:\.\d+)?)\s*%\s*soc", context, re.I)
            if path == "cell.r_dc_mohm":
                if tm:
                    self.add("cell.r_ref_temp_c", float(tm.group(1)), "°C", self._conf(False, "medium"), snippet, location)
                if sm:
                    self.add("cell.r_ref_soc_pct", float(sm.group(1)), "%", self._conf(False, "medium"), snippet, location)
            if d.conf_cap == "low":
                self.issue("DS_AMBIGUOUS_R", WARNING, path,
                           f"'{label}' has no AC/DC qualifier - stored as AC impedance, NOT used as DC resistance.",
                           "Confirm whether this is DCIR; if so enter it in the DC resistance field.")
        self.add(path, v, u, self._conf(assumed, d.conf_cap), snippet, location, note)

    # -- tables & blocks ---------------------------------------------------------------------------------
    def process_rows(self, rows: list[list[str]], where: str, title_hint: str = "") -> None:
        for title, block in split_blocks(rows):
            title = title or title_hint
            kind = classify_block(block)
            loc = f"{where} · {title}" if title else where
            if kind == "map":
                self._parse_map(block, title, loc)
            elif kind == "curve":
                self._parse_curve(block, title, loc)
            else:
                self._process_kv(block, loc)

    def _process_kv(self, block: list[list[str]], loc: str) -> None:
        """Key-value rows. If a header row names Value/Unit/Notes columns, columns are honoured by position
        (so an empty template cell is never confused with the notes text)."""
        hdr_idx, hdr, vname = 0, [], None
        for k, row in enumerate(block[:5]):                       # header may follow a title row
            cand = [c.strip().lower() for c in row]
            vname = next((n for n in ("value", "typical", "typ", "nominal", "values") if n in cand), None)
            if vname is not None:
                hdr_idx, hdr = k, cand
                break
        if vname is not None:
            vi = hdr.index(vname)
            ui = next((hdr.index(n) for n in ("unit", "units") if n in hdr), None)
            ni = next((hdr.index(n) for n in ("notes", "note", "condition", "conditions", "comment", "remarks") if n in hdr), None)
            for r in block[hdr_idx + 1:]:
                if not r or not r[0] or vi >= len(r) or not r[vi]:
                    continue
                unit = r[ui] if ui is not None and ui < len(r) else ""
                note = r[ni] if ni is not None and ni < len(r) else ""
                self.handle_pair(r[0], [r[vi], unit], loc, context=note)
            return
        for r in block:
            cells = [c for c in r if c != ""]
            if len(cells) >= 2:
                self.handle_pair(cells[0], cells[1:], loc)

    def _parse_curve(self, block: list[list[str]], title: str, loc: str) -> None:
        header, data = block[0], block[1:]
        if to_float(header[0]) is not None and to_float(header[1] if len(header) > 1 else "") is not None:
            header, data = ["", ""], block                                  # headerless numeric table
        pts = [(to_float(r[0]), to_float(r[1])) for r in data if len(r) >= 2]
        pts = [p for p in pts if p[0] is not None and p[1] is not None]
        if len(pts) < 2:
            return
        xh, yh = (header[0] or "").lower(), (header[1] or "").lower()
        t = re.sub(r"[_\-]+", " ", title.lower())
        txt = f"{t} {xh} {yh}"
        x_is_soc = "soc" in xh or ("soc" in t and "temp" not in xh)
        x_is_t = "temp" in xh or "°c" in xh or re.search(r"\bt\b", xh) or ("temp" in t and "soc" not in xh)
        key = None
        if x_is_soc and re.search(r"du\s?/\s?dt|dudt|entropic", txt):
            key = "dudt_vs_soc"
        elif x_is_soc and re.search(r"ocv|open.?circuit|\bvoltage", yh + " " + t):
            key = "ocv_vs_soc"
        elif x_is_soc and re.search(r"resist|ohm|dcir|dcr|\br\b|ω", yh + " " + t):
            key = "r_vs_soc"
        elif x_is_t and re.search(r"resist|ohm|dcir|dcr|\br\b|ω", yh + " " + t):
            key = "r_vs_temp"
        elif x_is_t and re.search(r"capacity|\bah\b|%", yh + " " + t):
            key = "capacity_vs_temp"
        if key is None:
            self.issue("DS_CURVE_UNRECOGNISED", INFO, loc, f"A numeric table ('{title or xh + '/' + yh}') could not be classified and was ignored.")
            return
        xs = [p[0] for p in pts]
        ys = [p[1] for p in pts]
        notes = []
        xl, yl = {"dudt_vs_soc": ("SOC [%]", "dU/dT [mV/K]"), "ocv_vs_soc": ("SOC [%]", "OCV [V]"),
                  "r_vs_soc": ("SOC [%]", "R [mΩ]"), "r_vs_temp": ("T [°C]", "R [mΩ]"),
                  "capacity_vs_temp": ("T [°C]", "Capacity [% of nominal]")}[key]
        if x_is_soc and max(xs) <= 1.0 and min(xs) >= 0.0 and ("fraction" in xh or "(-)" in xh or "[-]" in xh or max(xs) <= 1.0):
            xs = [x * 100 for x in xs]
            notes.append("SOC axis was a fraction (0-1): converted to %.")
        if key in ("r_vs_soc", "r_vs_temp"):
            cu = canon_unit(re.search(r"[\[(]\s*([^\])]+)[\])]", yh).group(1)) if re.search(r"[\[(]\s*([^\])]+)[\])]", yh) else ""
            f = UNITS["resistance"].get(cu)
            if f is None:
                notes.append("No resistance unit in the header: assumed mΩ.")
                f = 1.0
            ys = [y * f for y in ys]
        if key == "ocv_vs_soc" and re.search(r"\bmv\b", yh):
            ys = [y / 1000 for y in ys]
        if key == "dudt_vs_soc" and re.search(r"\[\s*v\s*/\s*k", yh):
            ys = [y * 1000 for y in ys]
        if key == "capacity_vs_temp":
            if re.search(r"\bah\b", yh):
                self.cap_curve_ah = {"x": xs, "y": ys, "loc": loc}
                return
            if max(ys) <= 1.5:
                ys = [y * 100 for y in ys]
                notes.append("Capacity given as a fraction: converted to %.")
        if len(set(xs)) != len(xs):
            self.issue("DS_CURVE_INVALID", ERROR, f"cell.{key}", f"{key}: duplicate x values in the table ({loc}) - curve ignored.")
            return
        order = sorted(range(len(xs)), key=lambda i: xs[i])
        self.res.curves[key] = {"x": [xs[i] for i in order], "y": [ys[i] for i in order], "x_label": xl, "y_label": yl,
                                "confidence": self.base_conf, "location": loc, "note": " ".join(notes)}

    def _parse_map(self, block: list[list[str]], title: str, loc: str) -> None:
        header, data = block[0], block[1:]
        ys = [to_float(c) for c in header[1:]]
        ys = [y for y in ys if y is not None]
        xs, z = [], []
        for r in data:
            x = to_float(r[0]) if r else None
            row = [to_float(c) for c in r[1:1 + len(ys)]]
            if x is None or len(row) < len(ys) or any(v is None for v in row):
                continue
            xs.append(x)
            z.append(row)
        if len(xs) < 2 or len(ys) < 2:
            return
        t = re.sub(r"[_]+", " ", (title + " " + header[0]).lower())
        key = "ocv_map" if re.search(r"ocv|open.?circuit|voltage", t) else "r_map" if re.search(r"resist|ohm|dcir|\br\b|ω", t) else None
        if key is None:
            self.issue("DS_CURVE_UNRECOGNISED", INFO, loc, "A 2-D table could not be classified (name it R_map or OCV_map) and was ignored.")
            return
        if any(b <= a for a, b in zip(xs, xs[1:])) or any(b <= a for a, b in zip(ys, ys[1:])):
            self.issue("DS_CURVE_INVALID", ERROR, f"cell.{key}", f"{key}: SOC or temperature axis is not strictly increasing ({loc}) - map ignored.")
            return
        note = ""
        if key == "r_map" and re.search(r"[\[(]\s*ω|[\[(]\s*ohm", t) and not re.search(r"m\s*(ω|ohm)", t):
            z = [[v * 1000 for v in r] for r in z]
            note = "Values were in Ω: converted to mΩ."
        self.res.maps[key] = {"x": xs, "y": ys, "z": z, "z_label": "R [mΩ]" if key == "r_map" else "OCV [V]",
                              "confidence": self.base_conf, "location": loc, "note": note}

    # -- text lines (PDF) --------------------------------------------------------------------------------
    def process_text_line(self, line: str, location: str) -> None:
        low = line.lower()
        hits = []
        for d, rx, ex in _COMPILED:
            for m in rx.finditer(low):
                if ex and ex.search(low[max(0, m.start() - 12):m.end() + 12]):
                    continue
                hits.append((m.start(), m.end(), d))
        hits.sort(key=lambda h: (h[0], -(h[1] - h[0])))
        keep, last_end = [], -1
        for s, e, d in hits:                                     # drop overlapping shorter matches
            if s >= last_end:
                keep.append((s, e, d))
                last_end = e
        for i, (s, e, d) in enumerate(keep):
            seg_end = keep[i + 1][0] if i + 1 < len(keep) else len(line)
            seg = line[e:seg_end].strip(" :=.-–\t")
            seg = re.sub(r"^[\[(][^\])]*[\])]\s*[:=]?\s*", "", seg)       # drop leading conditions such as "(1 kHz)"
            if not seg:
                continue
            label = line[s:e]
            self.handle_pair(label, [seg], location, context=line)


def split_blocks(rows: list[list[str]]) -> list[tuple[str, list[list[str]]]]:
    """Split rows into blocks separated by blank rows or ``[title]`` rows."""
    out: list[tuple[str, list[list[str]]]] = []
    cur: list[list[str]] = []
    title = ""

    def flush():
        nonlocal cur, title
        if cur:
            out.append((title, cur))
        cur = []

    for r in rows:
        nonempty = [c for c in r if c != ""]
        if not nonempty:
            flush()
            title = ""
            continue
        if len(nonempty) == 1 and nonempty[0].startswith("["):
            flush()
            title = nonempty[0].strip("[] ")
            continue
        if nonempty[0].startswith("#"):
            continue                                              # comment line
        cur.append(list(r))
    flush()
    return out


def _is_num(c: str) -> bool:
    return to_float(c) is not None


def classify_block(block: list[list[str]]) -> str:
    if len(block) < 3:
        return "kv"
    header, data = block[0], block[1:]
    hdr_ne = [c for c in header if c != ""]
    # 2-D map: header = [corner, T1, T2, ...] with numeric temperatures; first column numeric SOC
    if len(hdr_ne) >= 3 and sum(_is_num(c) for c in header[1:]) >= 2 and not _is_num(header[0]):
        if sum(_is_num(r[0]) for r in data if r) >= max(2, len(data) // 2):
            return "map"
    # curve: at least two numeric columns
    numeric_rows = sum(1 for r in data if len(r) >= 2 and _is_num(r[0]) and _is_num(r[1]))
    header_numeric = len(header) >= 2 and _is_num(header[0]) and _is_num(header[1])
    if numeric_rows >= max(2, int(0.8 * len(data))) and (not header_numeric or True):
        # a key-value block whose *value* column is numeric would also pass; distinguish by first column type
        first_col_numeric = sum(1 for r in data if r and _is_num(r[0]))
        if first_col_numeric >= int(0.8 * len(data)):
            return "curve"
    return "kv"


# ---------------------------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------------------------
def _pdf_backend():
    try:
        import pdfplumber
        return pdfplumber
    except (KeyboardInterrupt, SystemExit):
        raise
    except BaseException as exc:                                    # noqa: BLE001 - broken cryptography builds panic (pyo3)
        raise IngestionError("PDF support is unavailable in this environment (pdfplumber could not be imported: "
                             f"{type(exc).__name__}). Install a working 'cryptography'/'cffi' or upload CSV/Excel.") from exc


def _extract_pdf(ex: _Extractor, data: bytes) -> None:
    pdfplumber = _pdf_backend()
    try:
        pdf = pdfplumber.open(io.BytesIO(data))
    except Exception as exc:
        raise IngestionError(f"Could not open the PDF: {exc}") from exc
    n_chars = 0
    with pdf:
        for pno, page in enumerate(pdf.pages, 1):
            for tbl in page.extract_tables() or []:
                rows = [[(c or "").replace("\n", " ").strip() for c in r] for r in tbl]
                ex.process_rows(rows, f"page {pno} table")
            text = page.extract_text() or ""
            n_chars += len(text)
            for line in text.splitlines():
                if line.strip():
                    ex.process_text_line(line.strip(), f"page {pno}")
    if n_chars < 20:
        ex.issue("DS_PDF_NO_TEXT", ERROR, "", "The PDF contains no extractable text (scanned image?). OCR is not performed.",
                 "Upload a text-based PDF or an Excel/CSV datasheet, or enter the values manually.")


# ---------------------------------------------------------------------------------------------
# post-processing: A->C, missing list, suggestions
# ---------------------------------------------------------------------------------------------
REQUIRED = [
    ("cell.capacity_ah", "Nominal capacity", "Ah"), ("cell.v_nom", "Nominal voltage", "V"),
    ("cell.r_dc_mohm", "DC internal resistance (or R-vs-SOC/T table)", "mΩ"),
    ("cell.mass_kg", "Cell mass", "kg"), ("cell.cp_j_kg_k", "Specific heat capacity", "J/(kg·K)"),
]
RECOMMENDED = [
    ("cell.v_max", "Maximum voltage", "V"), ("cell.v_min", "Minimum voltage", "V"),
    ("cell.max_discharge_c", "Max continuous discharge C-rate", "C"), ("cell.max_charge_c", "Max continuous charge C-rate", "C"),
    ("cell.t_op_max_c", "Maximum operating temperature", "°C"), ("cell.t_op_min_c", "Minimum operating temperature", "°C"),
]
V_WINDOW = {"LFP": (3.65, 2.5), "NMC": (4.2, 3.0), "NCA": (4.2, 2.7), "LTO": (2.8, 1.5), "LCO": (4.2, 3.0), "LMO": (4.2, 3.0)}
WH_PER_KG = {"LFP": 160.0, "NMC": 240.0, "NCA": 250.0, "LTO": 90.0}


def _finalise(ex: _Extractor) -> DatasheetResult:
    res = ex.res
    seen: set[tuple] = set()
    uniq = []
    for i in res.issues:                                          # the same value is often seen in a table and in text
        key = (i.code, i.field, i.message)
        if key not in seen:
            seen.add(key)
            uniq.append(i)
    res.issues = uniq
    cap = res.fields.get("cell.capacity_ah")
    for path, (amps, snippet, loc) in ex.pending_amp.items():
        if cap:
            ex.add(path, amps / cap.value, "C", "medium", snippet, loc,
                   f"Converted from {amps:g} A using the extracted nominal capacity {cap.value:g} Ah - verify.")
        else:
            ex.issue("DS_CRATE_NEEDS_CAPACITY", WARNING, path, f"{path}: given as {amps:g} A but the nominal capacity is unknown, so it cannot be converted to C-rate.")
    if ex.cap_curve_ah is not None:
        if cap:
            ys = [y / cap.value * 100 for y in ex.cap_curve_ah["y"]]
            res.curves["capacity_vs_temp"] = {"x": ex.cap_curve_ah["x"], "y": ys, "x_label": "T [°C]", "y_label": "Capacity [% of nominal]",
                                              "confidence": ex.base_conf, "location": ex.cap_curve_ah["loc"], "note": "Converted from Ah to % of nominal."}
        else:
            ex.issue("DS_CURVE_NEEDS_CAPACITY", WARNING, "cell.capacity_vs_temp", "Capacity-vs-temperature is in Ah but nominal capacity is unknown.")

    f = res.fields
    chem = f["cell.chemistry"].value if "cell.chemistry" in f else None
    has_r = "cell.r_dc_mohm" in f or "r_vs_soc" in res.curves or "r_vs_temp" in res.curves or "r_map" in res.maps
    present = lambda p: p in f  # noqa: E731

    for path, label, unit in REQUIRED:
        ok = present(path) or (path == "cell.r_dc_mohm" and has_r)
        if ok:
            continue
        entry = {"path": path, "label": label, "unit": unit, "required": True, "suggestion": None}
        s = _suggest(path, f, chem)
        if s:
            res.suggestions[path] = s
            entry["suggestion"] = asdict(s)
        res.missing.append(entry)
        res.issues.append(Issue("DS_REQUIRED_MISSING", WARNING, path, f"Required parameter not found in the datasheet: {label}.",
                                "Enter it manually" + (" or review the offered engineering assumption." if s else ".")
                                + (" No default is offered for resistance: it must come from the datasheet or a test." if path == "cell.r_dc_mohm" else "")))
    for path, label, unit in RECOMMENDED:
        if not present(path):
            entry = {"path": path, "label": label, "unit": unit, "required": False, "suggestion": None}
            s = _suggest(path, f, chem)
            if s:
                res.suggestions[path] = s
                entry["suggestion"] = asdict(s)
            res.missing.append(entry)
    if "ocv_vs_soc" not in res.curves and "ocv_map" not in res.maps:
        res.missing.append({"path": "cell.ocv_vs_soc", "label": "OCV vs SOC curve", "unit": "", "required": False, "suggestion": None,
                            "note": "Without it OCV is held at the nominal voltage."})
    if "dudt_vs_soc" not in res.curves and "entropic.constant_mv_per_k" not in f and "ocv_map" not in res.maps:
        res.missing.append({"path": "cell.dudt_vs_soc", "label": "Entropic coefficient dU/dT", "unit": "mV/K", "required": False, "suggestion": None,
                            "note": "Reversible heat cannot be calculated accurately; enter an estimate or exclude it explicitly."})
    if not f and not res.curves and not res.maps:
        ex.issue("DS_NOTHING_FOUND", ERROR, "", "No cell parameters could be recognised in this file.",
                 "Use the datasheet template (Parameter / Value / Unit) or enter values manually.")
    return res


def _suggest(path: str, f: dict, chem: str | None) -> Suggestion | None:
    if path == "cell.cp_j_kg_k":
        return Suggestion(path, 1000.0, "J/(kg·K)", "low",
                          "Typical Li-ion cell specific heat 800-1100 J/(kg·K). Not a datasheet value - measure or confirm with the cell maker.")
    if path == "cell.mass_kg" and chem in WH_PER_KG and "cell.capacity_ah" in f and "cell.v_nom" in f:
        e_wh = f["cell.capacity_ah"].value * f["cell.v_nom"].value
        return Suggestion(path, round(e_wh / WH_PER_KG[chem] , 3), "kg", "low",
                          f"Rough estimate from a typical {chem} cell-level specific energy of {WH_PER_KG[chem]:g} Wh/kg (±25 %). Enter the real cell mass.")
    if path in ("cell.v_max", "cell.v_min") and chem in V_WINDOW:
        v = V_WINDOW[chem][0 if path == "cell.v_max" else 1]
        return Suggestion(path, v, "V", "low", f"Typical {chem} voltage window - confirm against the cell datasheet.")
    return None


# ---------------------------------------------------------------------------------------------
# public API
# ---------------------------------------------------------------------------------------------
def parse_datasheet(data: bytes, filename: str) -> DatasheetResult:
    kind = file_kind(filename)
    if kind == "pdf":
        ex = _Extractor(filename, "pdf")
        _extract_pdf(ex, data)
        return _finalise(ex)
    sheets: list[SheetRows] = read_sheets(data, filename)
    ex = _Extractor(filename, kind)
    for sh in sheets:
        ex.process_rows(sh.rows, f"sheet '{sh.name}'" if kind == "xlsx" else filename,
                        title_hint=sh.name if kind == "xlsx" else "")
    return _finalise(ex)


# ---------------------------------------------------------------------------------------------
# templates
# ---------------------------------------------------------------------------------------------
TEMPLATE_KV = [
    ("Cell name / part number", "", "", ""), ("Chemistry", "", "", "LFP / NMC / NCA / LTO ..."), ("Cell type", "", "", "prismatic / cylindrical / pouch"),
    ("Nominal capacity", "", "Ah", ""), ("Nominal voltage", "", "V", ""), ("Maximum voltage", "", "V", "charge cut-off"),
    ("Minimum voltage", "", "V", "discharge cut-off"), ("DC internal resistance (DCIR)", "", "mΩ", "state SOC / temperature / pulse length"),
    ("AC impedance (1 kHz)", "", "mΩ", "information only"), ("Max continuous discharge current", "", "C", "or A"),
    ("Max continuous charge current", "", "C", "or A"), ("Pulse discharge current", "", "C", "or A; note pulse duration"),
    ("Pulse duration", "", "s", ""), ("Length", "", "mm", ""), ("Width", "", "mm", ""), ("Height", "", "mm", ""),
    ("Diameter", "", "mm", "cylindrical only"), ("Cell mass", "", "kg", ""), ("Specific heat capacity", "", "J/(kg·K)", "rarely on datasheets"),
    ("Operating temperature (discharge)", "", "°C", "min ~ max"), ("Operating temperature (charge)", "", "°C", "min ~ max"),
    ("Recommended temperature", "", "°C", "min ~ max"),
]
TEMPLATE_TABLES = [
    ("R_vs_SOC", ["SOC (%)", "R_DC (mΩ)"], 5), ("R_vs_T", ["Temperature (°C)", "R_DC (mΩ)"], 5),
    ("OCV_vs_SOC", ["SOC (%)", "OCV (V)"], 6), ("dUdT_vs_SOC", ["SOC (%)", "dU/dT (mV/K)"], 4),
    ("Capacity_vs_T", ["Temperature (°C)", "Capacity (% of nominal)"], 5),
]


def datasheet_template_csv() -> str:
    lines = ["# Cell datasheet template - fill Value/Unit, delete rows you do not have. Curves below: one x,y pair per row.",
             "Parameter,Value,Unit,Notes"]
    for p, v, u, n in TEMPLATE_KV:
        lines.append(f"{p},{v},{u},{n}")
    for name, hdr, n in TEMPLATE_TABLES:
        lines += ["", f"[{name}]", ",".join(hdr)] + [","] * n
    lines += ["", "[R_map]", "SOC (%) \\ T (°C),-10,0,25,40"] + ["0,,,,", "50,,,,", "100,,,,"]
    return "\n".join(lines) + "\n"


def datasheet_template_xlsx() -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Cell"
    ws.append(["Parameter", "Value", "Unit", "Notes"])
    for row in TEMPLATE_KV:
        ws.append(list(row))
    for name, hdr, n in TEMPLATE_TABLES:
        s = wb.create_sheet(name)
        s.append(hdr)
        for _ in range(n):
            s.append([None, None])
    m = wb.create_sheet("R_map")
    m.append(["SOC (%) \\ T (°C)", -10, 0, 25, 40])
    for soc in (0, 50, 100):
        m.append([soc, None, None, None, None])
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()
